"""
Tests for Audit Trail (hash chain) — Excel logger + validator.
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from agents.excel_logger import ExcelLogger, GENESIS_HASH, SHEET_SCHEMAS
from utils.audit import AuditValidator


class TestExcelLogger:

    @pytest.fixture
    def temp_xlsx(self):
        """Provide a temporary xlsx file path."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            path = Path(tmp.name)
        # Clean up before tests (start fresh)
        if path.exists():
            path.unlink()
        yield path
        if path.exists():
            path.unlink()

    def test_init_creates_workbook_with_sheets(self, temp_xlsx):
        """Initialization should create 6 sheets with headers."""
        logger = ExcelLogger(xlsx_path=temp_xlsx)
        assert temp_xlsx.exists()
        assert logger.wb is not None
        for sheet_name in SHEET_SCHEMAS.keys():
            assert sheet_name in logger.wb.sheetnames, f"Sheet '{sheet_name}' missing"

    def test_log_creates_entry_with_hash_chain(self, temp_xlsx):
        """Logging should add a row with prev_hash + this_hash."""
        logger = ExcelLogger(xlsx_path=temp_xlsx)
        logger.log("Audit", {"event_type": "test", "details": "test event"})
        logger.close()

        # Verify with openpyxl
        try:
            from openpyxl import load_workbook
            wb = load_workbook(temp_xlsx)
            ws = wb["Audit"]
            # Row 2 should have our event
            assert ws.cell(row=2, column=1).value is not None  # timestamp
            assert ws.cell(row=2, column=2).value == "test"  # event_type
            assert ws.cell(row=2, column=4).value == GENESIS_HASH  # prev_hash
            assert ws.cell(row=2, column=5).value is not None  # this_hash
        except ImportError:
            pass  # openpyxl not installed — skip verification

    def test_hash_chain_progresses(self, temp_xlsx):
        """Second event's prev_hash should equal first event's this_hash."""
        logger = ExcelLogger(xlsx_path=temp_xlsx)
        logger.log("Audit", {"event_type": "first", "details": "event 1"})
        first_this_hash = logger.last_hash
        logger.log("Audit", {"event_type": "second", "details": "event 2"})
        second_prev_hash = logger.last_hash  # Actually this is the second's this_hash

        # The second event's prev_hash should equal first's this_hash
        # (We can verify by checking the Excel file)
        try:
            from openpyxl import load_workbook
            wb = load_workbook(temp_xlsx)
            ws = wb["Audit"]
            # Row 3 = second event
            second_prev = ws.cell(row=3, column=4).value
            assert second_prev == first_this_hash, \
                f"Chain broken: row 3 prev_hash={second_prev} != row 2 this_hash={first_this_hash}"
        except ImportError:
            pass

    def test_verify_chain_intact(self, temp_xlsx):
        """Verify chain returns valid=True for unmodified log."""
        logger = ExcelLogger(xlsx_path=temp_xlsx)
        logger.log("Audit", {"event_type": "e1", "details": "d1"})
        logger.log("Audit", {"event_type": "e2", "details": "d2"})
        logger.log("Audit", {"event_type": "e3", "details": "d3"})
        logger.close()

        validator = AuditValidator(temp_xlsx)
        result = validator.verify_all_sheets()
        assert result["overall_valid"] is True
        assert result["total_rows"] >= 3


class TestAuditValidator:

    @pytest.fixture
    def populated_xlsx(self, temp_xlsx=None):
        """Create an xlsx with some events."""
        if temp_xlsx is None:
            with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
                temp_xlsx = Path(tmp.name)
            if temp_xlsx.exists():
                temp_xlsx.unlink()

        logger = ExcelLogger(xlsx_path=temp_xlsx)
        logger.log("Audit", {"event_type": "e1", "details": "first"})
        logger.log("Audit", {"event_type": "e2", "details": "second"})
        logger.log("Process", {"step": "INGEST", "status": "ok", "agent": "ingest", "duration_sec": 1.5})
        logger.close()

        yield temp_xlsx

        if temp_xlsx.exists():
            temp_xlsx.unlink()

    def test_verify_returns_markdown_report(self, populated_xlsx):
        """generate_audit_report should return a markdown string."""
        validator = AuditValidator(populated_xlsx)
        report = validator.generate_audit_report()
        assert isinstance(report, str)
        assert "AUDIT TRAIL INTEGRITY REPORT" in report
        assert "VALID" in report or "TAMPERED" in report

    def test_verify_nonexistent_file(self):
        """Verify should return error for missing file."""
        validator = AuditValidator(Path("/nonexistent/file.xlsx"))
        result = validator.verify_all_sheets()
        assert result["overall_valid"] is False
        assert "error" in result
